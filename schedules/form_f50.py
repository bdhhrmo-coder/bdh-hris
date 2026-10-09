"""
Printable Personnel Work Schedule - Form BDH-ADM-AO-01F50, Revision 2
(Batch 5, owner decisions 2026-10-09).

- Long bond (8.5 x 13 in), landscape, the WHOLE MONTH ON ONE PAGE: Name,
  Designation, one narrow column per day, Total Hours, Total Days.
- Paperless: instead of signature lines, five "digitally ..." stamps from
  the schedule's action log - Prepared, Reviewed (HR), Recommending
  Approval (AO), Approved (COH) and Recorded (HR) - with the real name,
  position, date and time. Steps not yet reached print PENDING; a returned
  schedule prints the red RETURNED banner with the remark.
- Only the current cycle (since the latest submit/resubmit) is stamped, so
  a resubmitted schedule restarts its stamps; earlier cycles stay in the
  on-screen history.
- Cells changed after approval (approved exchange / HR correction) print
  their current value; the changes are listed under the grid.
"""

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from django.utils import timezone

from leave.pdf_convert import workbook_to_pdf
from printouts.sheet import BDH_LOGO, DONE_COLOR, FOLIO, PALAWAN_SEAL, PENDING_COLOR, STOPPED_COLOR
from printouts.stamps import StepDef, actor_name, compute_stamps

from . import rules
from .models import ScheduleCell, ShiftCode

FORM_CODE = "BDH-ADM-AO-01F50"
REVISION = 2
FONT = "Arial"
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="DDDDDD")
WEEKEND_FILL = PatternFill("solid", fgColor="F2F2F2")
STAMP_HEAD_FILL = PatternFill("solid", fgColor="DDDDDD")

ROUTING = [
    StepDef("Prepared", ("submit", "resubmit"), "DIGITALLY PREPARED"),
    StepDef("Reviewed by (HR)", "review", "DIGITALLY REVIEWED"),
    StepDef("Recommending Approval (AO)", "recommend", "DIGITALLY RECOMMENDED"),
    StepDef("Approved (Chief of Hospital)", "approve", "DIGITALLY APPROVED"),
    StepDef("Recorded by HR", "record", "DIGITALLY RECORDED"),
]

NAME_COL, DESIG_COL, FIRST_DAY_COL = 1, 2, 3  # same columns as the paper F50


def reference_no(schedule):
    return f"PWS-{schedule.pk:06d}"


def current_cycle_actions(schedule):
    actions = list(schedule.actions.select_related("acted_by__employee").order_by("acted_at", "pk"))
    starts = [i for i, a in enumerate(actions) if a.action in ("submit", "resubmit")]
    return actions[starts[-1]:] if starts else []


def _put(ws, row, c1, c2, value, *, bold=False, italic=False, size=8, align="center", color=None, fill=None,
         border=None, wrap=True):
    if c2 > c1:
        ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    cell = ws.cell(row=row, column=c1, value=value)
    cell.font = Font(name=FONT, bold=bold, italic=italic, size=size, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    for c in range(c1, c2 + 1):
        if fill:
            ws.cell(row=row, column=c).fill = fill
        if border:
            ws.cell(row=row, column=c).border = border
    return cell


def _outline(ws, top, bottom, c1, c2, color, style):
    """Rectangle around rows top..bottom, columns c1..c2 (numbers)."""
    side = Side(style=style, color=color)
    for r in range(top, bottom + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            b = cell.border
            cell.border = Border(top=side if r == top else b.top, bottom=side if r == bottom else b.bottom,
                                 left=side if c == c1 else b.left, right=side if c == c2 else b.right)


def _spans(first, last, parts):
    """Split columns first..last into `parts` near-equal (c1, c2) spans."""
    n = last - first + 1
    out, start = [], first
    for i in range(parts):
        width = n // parts + (1 if i < n % parts else 0)
        out.append((start, start + width - 1))
        start += width
    return out


def fill_schedule_form(schedule, printed_by=None):
    days = schedule.days
    last_day_col = FIRST_DAY_COL + len(days) - 1
    hours_col, total_days_col = last_day_col + 1, last_day_col + 2
    last_col = total_days_col
    L = get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "F50"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions[L(NAME_COL)].width = 24
    ws.column_dimensions[L(DESIG_COL)].width = 13
    for c in range(FIRST_DAY_COL, last_day_col + 1):
        ws.column_dimensions[L(c)].width = 4.6
    ws.column_dimensions[L(hours_col)].width = 6.5
    ws.column_dimensions[L(total_days_col)].width = 6.5

    # -- header ---------------------------------------------------------------
    for path, anchor in ((PALAWAN_SEAL, "A1"), (BDH_LOGO, "B1")):
        if path.exists():
            img = XLImage(str(path))
            img.width = img.height = 48
            ws.add_image(img, anchor)
    _put(ws, 1, last_col - 7, last_col, f"Form No. {FORM_CODE}", italic=True, size=8, align="right")
    _put(ws, 2, last_col - 7, last_col, f"Revision: {REVISION}", italic=True, size=8, align="right")
    _put(ws, 1, FIRST_DAY_COL, last_col - 8, "Republic of the Philippines · PROVINCIAL GOVERNMENT OF PALAWAN",
         size=8, italic=True)
    _put(ws, 2, FIRST_DAY_COL, last_col - 8, "BATARAZA DISTRICT HOSPITAL", bold=True, size=11)
    _put(ws, 3, FIRST_DAY_COL, last_col - 8, "PERSONNEL WORK SCHEDULE", bold=True, size=13)
    ws.row_dimensions[3].height = 18
    _put(ws, 5, 1, DESIG_COL + 9, f"Department/Section/Unit:  {schedule.area_full_name}", bold=True, size=9,
         align="left")
    _put(ws, 5, DESIG_COL + 10, last_col, f"Month/Year:  {schedule.month_label}", bold=True, size=9, align="right")

    # -- grid header (2 rows: day number, weekday) ---------------------------------
    h1, h2 = 7, 8
    for col, label in ((NAME_COL, "NAME\n(Last name, First name)"), (DESIG_COL, "DESIGNATION"),
                       (hours_col, "TOTAL HOURS"), (total_days_col, "TOTAL DAYS")):
        ws.merge_cells(start_row=h1, start_column=col, end_row=h2, end_column=col)
        _put(ws, h1, col, col, label, bold=True, size=7)
    for i, d in enumerate(days):
        col = FIRST_DAY_COL + i
        _put(ws, h1, col, col, d.day, bold=True, size=8)
        _put(ws, h2, col, col, d.strftime("%a")[:2], size=7)
    for r in (h1, h2):
        for c in range(1, last_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill, cell.border = HEAD_FILL, BOX

    # -- rows ----------------------------------------------------------------------------
    cells = {}
    for c in ScheduleCell.objects.filter(row__schedule=schedule).select_related("shift"):
        cells[(c.row_id, c.date)] = c
    r = h2 + 1
    rows = list(schedule.rows.select_related("employee"))
    for row in rows:
        row_cells = [cells.get((row.pk, d)) for d in days]
        hours, duty_days = rules.row_totals([c for c in row_cells if c])
        _put(ws, r, NAME_COL, NAME_COL, f"{row.employee.surname.upper()}, {row.employee.first_name}", size=8,
             align="left", wrap=False)
        _put(ws, r, DESIG_COL, DESIG_COL, row.designation, size=7, align="left", wrap=False)
        for i, (d, c) in enumerate(zip(days, row_cells)):
            col = FIRST_DAY_COL + i
            cell = _put(ws, r, col, col, c.shift.code if c and c.shift_id else "", size=6.5)
            if d.weekday() >= 5:
                cell.fill = WEEKEND_FILL
        _put(ws, r, hours_col, hours_col, f"{hours.normalize():f}", bold=True, size=8)
        _put(ws, r, total_days_col, total_days_col, duty_days, bold=True, size=8)
        for c in range(1, last_col + 1):
            ws.cell(row=r, column=c).border = BOX
        ws.row_dimensions[r].height = 14
        r += 1
    if not rows:
        _put(ws, r, 1, last_col, "No employees on this schedule.", italic=True, size=8, border=BOX)
        r += 1

    # -- legend, notes, changes ---------------------------------------------------------
    r += 1
    legend = "   ".join(
        f"{c.code} = {c.description or ''}{' (' + format(c.paid_hours, 'g') + ' h)' if c.is_duty else ''}".strip()
        for c in ShiftCode.objects.filter(is_active=True))
    _put(ws, r, 1, last_col, f"Shift codes:  {legend}", size=7, align="left")
    r += 1
    _put(ws, r, 1, last_col, f"NOTES/ REMARKS:  {schedule.notes or ''}", size=8, align="left")
    r += 1
    changes = list(schedule.changes.select_related("employee"))
    if changes:
        shown = changes[:6]
        text = "; ".join(f"{ch.employee.surname.upper()} {ch.date:%b %d}: {ch.old_shift or 'blank'} → "
                         f"{ch.new_shift or 'blank'} ({ch.reason})" for ch in shown)
        if len(changes) > len(shown):
            text += f"; and {len(changes) - len(shown)} more (see the HRIS)"
        _put(ws, r, 1, last_col, f"Changed after approval:  {text}", size=7, italic=True, align="left")
        ws.row_dimensions[r].height = 22 if len(text) > 220 else 12
        r += 1

    # -- stamps: one row of five -------------------------------------------------------------
    r += 1
    stamps, stopped = compute_stamps(current_cycle_actions(schedule), ROUTING)
    for stamp, (c1, c2) in zip(stamps, _spans(1, last_col, len(stamps))):
        done = stamp.state == "done"
        color = DONE_COLOR if done else PENDING_COLOR
        _put(ws, r, c1, c2, stamp.role_label.upper(), bold=True, size=7, color="333333", fill=STAMP_HEAD_FILL)
        _put(ws, r + 1, c1, c2, stamp.text, bold=done, italic=not done, size=8, color=color)
        _put(ws, r + 2, c1, c2, stamp.when, bold=done, size=7, color=color)
        _put(ws, r + 3, c1, c2, stamp.name or "—", bold=True, size=8)
        _put(ws, r + 4, c1, c2, stamp.position, italic=True, size=7)
        _outline(ws, r, r + 4, c1, c2, color if done else "999999", "medium" if done else "thin")
    r += 6
    if stopped is not None:
        text = f"✖ {stopped.text}  by  {stopped.name}"
        if stopped.position:
            text += f", {stopped.position}"
        text += f"  —  {stopped.when}"
        if stopped.notes:
            text += f"   ·   Remarks: {stopped.notes}"
        _put(ws, r, 1, last_col, text, bold=True, size=8, color=STOPPED_COLOR)
        _outline(ws, r, r, 1, last_col, STOPPED_COLOR, "medium")
        r += 2

    printed = f"Printed {timezone.localtime():%b %d, %Y %I:%M %p}"
    if printed_by is not None:
        printed += f" by {actor_name(printed_by)}"
    _put(ws, r, 1, last_col,
         "Processed electronically through the BDH HRIS. Each stamp records the user account, date and time of "
         f"that action from the system's action log.     Reference No.: {reference_no(schedule)}     ·     {printed}",
         italic=True, size=7, align="left")

    ps = ws.page_setup
    ps.orientation = "landscape"
    ps.paperSize = FOLIO
    ps.fitToWidth = 1
    ps.fitToHeight = 1   # whole month on ONE page (owner decision)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.35
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:{L(last_col)}{r}"
    return wb


def render_pdf(schedule, printed_by=None) -> bytes:
    return workbook_to_pdf(fill_schedule_form(schedule, printed_by), "Personnel Work Schedule")
