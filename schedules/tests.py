"""Batch 5: duty schedules (F50) and the Exchange of Duty schedule checks."""

import io
import re
import shutil
from datetime import date
from unittest import skipUnless

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from accounts.models import RoleAssignment
from auditlog.aggregation import audit_rows
from employees.models import Employee
from homepage.kpis import waiting_for_me
from notifications.models import Notification
from orgstructure.models import Section

from . import rules, services
from .models import DutySchedule, ScheduleCell, ScheduleChange, ShiftCode

S = DutySchedule
YEAR, MONTH = 2027, 3  # a fixed future month: March 2027 (31 days)


def emp(username, section=None, role=None, role_section=None, position="Nurse II"):
    user = User.objects.create_user(username, password="x")
    e = Employee.objects.create(user=user, employee_id=username.upper(), surname=username.title(), first_name="T",
                                position=position, date_hired=date(2024, 1, 1), employment_status="REGULAR")
    if section:
        e.sections.add(section)
    if role:
        RoleAssignment.objects.create(employee=e, role=role, section=role_section)
    return e


def code(c):
    return ShiftCode.objects.get(code=c)


def set_days(schedule, employee, mapping):
    """mapping: {day number: code or None}"""
    row = schedule.rows.get(employee=employee)
    for day, c in mapping.items():
        ScheduleCell.objects.filter(row=row, date=date(schedule.year, schedule.month, day)).update(
            shift=code(c) if c else None)


class Base(TestCase):
    def setUp(self):
        self.lab = Section.objects.get(name="Laboratory Section")
        self.pharm = Section.objects.get(name="Pharmacy Section")
        self.sup = emp("schsup", role=RoleAssignment.SUPERVISOR, role_section=self.lab, position="Medical Technologist III")
        self.hr = emp("schhr", role=RoleAssignment.HR_PROCESSOR, position="HR Assistant")
        self.hr2 = emp("schhr2", role=RoleAssignment.HR_ADMINISTRATOR, position="HRMO")
        self.ao = emp("schao", role=RoleAssignment.ADMINISTRATIVE_OFFICER, position="AO V")
        self.coh = emp("schcoh", role=RoleAssignment.CHIEF_OF_HOSPITAL, position="Chief of Hospital")
        self.alice = emp("schalice", self.lab, position="Medical Technologist I")
        self.bob = emp("schbob", self.lab, position="Medical Technologist I")
        self.carl = emp("schcarl", self.pharm, position="Pharmacist I")
        self.outsider = emp("schout", self.pharm)

    def new_schedule(self, preparer=None, section=None, year=YEAR, month=MONTH):
        preparer = preparer or self.sup
        section = section or self.lab
        role = rules.preparer_role(preparer, section=section)
        return services.create_schedule(preparer.user, role, year, month, section=section)

    def post_action(self, who, schedule, action, notes=""):
        self.client.force_login(who.user)
        return self.client.post(reverse("schedules:action", args=[schedule.pk]), {"action": action, "notes": notes})

    def route_to(self, schedule, status):
        if schedule.is_editable:
            services.submit(schedule, schedule.prepared_by)
        for step_status, actor in ((S.REVIEWED, self.hr), (S.RECOMMENDED, self.ao), (S.APPROVED, self.coh),
                                   (S.RECORDED, self.hr)):
            if schedule.status == status:
                return schedule
            action = rules.STEPS[schedule.status][1]
            services.act(schedule, actor.user, action)
        return schedule


class CreateAndEditTests(Base):
    def test_create_adds_area_employees_and_default_cutoff(self):
        s = self.new_schedule()
        self.assertEqual(set(s.rows.values_list("employee", flat=True)),
                         {self.alice.pk, self.bob.pk})
        self.assertEqual(s.preparer_role, S.FILER_SUPERVISOR)
        self.assertEqual(s.cutoff_date, date(2027, 2, 20))
        self.assertEqual(s.rows.first().cells.count(), 31)

    def test_new_schedule_page_creates_and_reopens(self):
        self.client.force_login(self.sup.user)
        page = self.client.get(reverse("schedules:new"))
        self.assertContains(page, "Laboratory Section")
        self.assertNotContains(page, "Pharmacy Section")
        data = {"area": f"s{self.lab.pk}", "month": MONTH, "year": YEAR}
        r = self.client.post(reverse("schedules:new"), data)
        s = DutySchedule.objects.get()
        self.assertRedirects(r, reverse("schedules:edit", args=[s.pk]), fetch_redirect_response=False)
        self.assertEqual(self.client.get(reverse("schedules:edit", args=[s.pk])).status_code, 200)
        r = self.client.post(reverse("schedules:new"), data)  # same month again -> opens the existing one
        self.assertEqual(DutySchedule.objects.count(), 1)
        for name in ("list", "mine"):
            self.assertEqual(self.client.get(reverse(f"schedules:{name}")).status_code, 200)
        self.assertEqual(self.client.get(reverse("schedules:detail", args=[s.pk])).status_code, 200)

    def test_shift_codes_are_hr_administrator_only(self):
        self.client.force_login(self.hr.user)
        self.assertEqual(self.client.get(reverse("schedules:shift_codes")).status_code, 403)
        self.client.force_login(self.hr2.user)
        self.assertContains(self.client.get(reverse("schedules:shift_codes")), "7A-7P")
        self.client.post(reverse("schedules:shift_code_new"), {"code": "4-12", "description": "Evening",
                                                               "paid_hours": "8", "is_duty": "on", "is_active": "on",
                                                               "sort_order": 5})
        self.assertTrue(ShiftCode.objects.filter(code="4-12", paid_hours=8).exists())

    def test_only_area_supervisor_or_hr_can_prepare(self):
        self.client.force_login(self.alice.user)
        self.assertEqual(self.client.get(reverse("schedules:new")).status_code, 403)
        self.assertEqual(rules.preparer_role(self.sup, section=self.pharm), None)
        self.assertEqual(rules.preparer_role(self.hr, section=self.pharm), S.FILER_HR)

    def test_grid_post_saves_cells_and_totals(self):
        s = self.new_schedule()
        row = s.rows.get(employee=self.alice)
        twelve = code("7A-7P").pk
        self.client.force_login(self.sup.user)
        data = {f"c-{row.pk}-{d}": twelve for d in (1, 2, 3)}
        data.update({f"c-{row.pk}-{d}": code("OFF").pk for d in (4, 5, 6)})
        data["notes"] = "Lunch break hours 12:00NN-1:00PM"
        self.client.post(reverse("schedules:edit", args=[s.pk]), data)
        hours, days = rules.row_totals(list(row.cells.select_related("shift")))
        self.assertEqual((hours, days), (36, 3))  # 3 x 12 h, OFF counts for nothing

    def test_eight_to_five_is_eight_paid_hours(self):
        self.assertEqual(code("8-5").paid_hours, 8)
        self.assertTrue(code("7P-7A").is_duty)
        self.assertFalse(code("OFF").is_duty)

    def test_double_booking_is_flagged_not_blocked(self):
        s1 = self.new_schedule()
        s2 = self.new_schedule(preparer=self.hr, section=self.pharm)
        services.add_row(s2, self.alice)
        set_days(s1, self.alice, {5: "7A-7P"})
        set_days(s2, self.alice, {5: "7P-7A"})
        conflicts = services.double_bookings(s1)
        self.assertEqual([(e, d) for e, d, _ in conflicts], [(self.alice, date(YEAR, MONTH, 5))])
        services.submit(s1, s1.prepared_by)  # still allowed
        self.assertEqual(s1.status, S.SUBMITTED)


class RoutingTests(Base):
    def test_full_route_locks_then_publishes(self):
        s = self.new_schedule()
        set_days(s, self.alice, {1: "8-5"})
        services.submit(s, self.sup.user)
        self.assertEqual(self.post_action(self.hr, s, "review").status_code, 302)
        self.post_action(self.ao, s, "recommend")
        self.post_action(self.coh, s, "approve")
        s.refresh_from_db()
        self.assertEqual(s.status, S.APPROVED)
        self.assertIsNotNone(s.approved_at)
        self.assertEqual(s.approved_by, self.coh.user)
        self.assertTrue(s.is_locked)
        with self.assertRaises(ValidationError):
            services.save_grid(s, {})
        # not on My Schedule until HR records it
        self.client.force_login(self.alice.user)
        page = self.client.get(reverse("schedules:mine") + f"?y={YEAR}&m={MONTH}")
        self.assertNotContains(page, "Laboratory Section")
        self.post_action(self.hr2, s, "record")
        s.refresh_from_db()
        self.assertEqual(s.status, S.RECORDED)
        self.client.force_login(self.alice.user)
        page = self.client.get(reverse("schedules:mine") + f"?y={YEAR}&m={MONTH}")
        self.assertContains(page, "Laboratory Section")
        self.assertTrue(Notification.objects.filter(recipient=self.alice, message__contains="duty schedule").exists())

    def test_steps_must_be_in_order_and_by_the_right_role(self):
        s = self.new_schedule()
        services.submit(s, self.sup.user)
        self.assertEqual(self.post_action(self.ao, s, "recommend").status_code, 403)
        self.assertEqual(self.post_action(self.alice, s, "review").status_code, 403)
        self.assertEqual(self.post_action(self.coh, s, "approve").status_code, 403)

    def test_preparer_never_acts_later(self):
        s = self.new_schedule(preparer=self.hr)  # HR prepares
        self.assertEqual(s.preparer_role, S.FILER_HR)
        services.submit(s, self.hr.user)
        self.assertEqual(self.post_action(self.hr, s, "review").status_code, 403)
        self.post_action(self.hr2, s, "review")  # a different HR person reviews
        s.refresh_from_db()
        self.assertEqual(s.status, S.REVIEWED)
        services.act(s, self.ao.user, "recommend")
        services.act(s, self.coh.user, "approve")
        self.assertEqual(self.post_action(self.hr, s, "record").status_code, 403)

    def test_return_needs_remark_and_resubmit_restarts(self):
        s = self.new_schedule()
        services.submit(s, self.sup.user)
        services.act(s, self.hr.user, "review")
        self.post_action(self.ao, s, "return", "")
        s.refresh_from_db()
        self.assertEqual(s.status, S.REVIEWED)  # no remark -> not returned
        self.post_action(self.ao, s, "return", "Day 14 has no nurse on night duty.")
        s.refresh_from_db()
        self.assertEqual(s.status, S.RETURNED)
        self.assertTrue(s.is_editable)
        self.client.force_login(self.sup.user)
        self.client.post(reverse("schedules:edit", args=[s.pk]), {"then": "submit", "notes": s.notes})
        s.refresh_from_db()
        self.assertEqual(s.status, S.SUBMITTED)  # back to HR review
        from .form_f50 import current_cycle_actions
        self.assertEqual([a.action for a in current_cycle_actions(s)], ["resubmit"])

    def test_waiting_for_me_and_homepage(self):
        s = self.new_schedule()
        services.submit(s, self.sup.user)
        self.assertIn(("Duty schedule", 1, "schedules:list"), waiting_for_me(self.hr, self.hr.user))
        self.assertEqual([i for i in waiting_for_me(self.ao, self.ao.user) if i[0] == "Duty schedule"], [])

    def test_visibility(self):
        s = self.new_schedule()
        self.client.force_login(self.alice.user)
        self.assertEqual(self.client.get(reverse("schedules:detail", args=[s.pk])).status_code, 403)
        self.client.force_login(self.outsider.user)
        self.assertEqual(self.client.get(reverse("schedules:detail", args=[s.pk])).status_code, 403)
        self.client.force_login(self.ao.user)
        self.assertEqual(self.client.get(reverse("schedules:detail", args=[s.pk])).status_code, 200)

    def test_audit_log_has_steps(self):
        s = self.new_schedule()
        self.route_to(s, S.REVIEWED)
        actions = [r["action"] for r in audit_rows(module="Duty Schedule")]
        self.assertEqual(sorted(actions), ["create", "review", "submit"])


class AfterApprovalTests(Base):
    def test_hr_correction_is_logged_and_notified(self):
        s = self.new_schedule()
        set_days(s, self.alice, {10: "7A-7P"})
        self.route_to(s, S.RECORDED)
        row = s.rows.get(employee=self.alice)
        self.client.force_login(self.sup.user)
        r = self.client.post(reverse("schedules:correct", args=[s.pk]),
                             {"row": row.pk, "date": "2027-03-10", "shift": code("OFF").pk, "reason": "x"})
        self.assertEqual(r.status_code, 403)  # Supervisor can't correct
        self.client.force_login(self.hr.user)
        self.client.post(reverse("schedules:correct", args=[s.pk]),
                         {"row": row.pk, "date": "2027-03-10", "shift": code("OFF").pk, "reason": ""})
        self.assertFalse(ScheduleChange.objects.exists())  # reason required
        self.client.post(reverse("schedules:correct", args=[s.pk]),
                         {"row": row.pk, "date": "2027-03-10", "shift": code("OFF").pk, "reason": "Sick leave"})
        ch = ScheduleChange.objects.get()
        self.assertEqual((ch.old_shift, ch.new_shift), ("7A-7P", "OFF"))
        self.assertIn("change", [r["action"] for r in audit_rows(module="Duty Schedule", employee_id=self.alice.pk)])
        self.assertTrue(Notification.objects.filter(recipient=self.alice, message__contains="Sick leave").exists())


class ExcelTests(Base):
    def test_round_trip(self):
        from openpyxl import load_workbook

        from .excel import build_template

        s = self.new_schedule()
        wb = build_template(s)
        ws = wb["Schedule"]
        # find alice's row and fill days 1-3 with 7A-7P, day 4 with a bad code
        for r in range(1, ws.max_row + 1):
            if ws.cell(row=r, column=1).value == self.alice.employee_id:
                ws.cell(row=r, column=4, value="7A-7P")
                ws.cell(row=r, column=5, value="7a-7p")
                ws.cell(row=r, column=6, value="7A-7P")
                ws.cell(row=r, column=7, value="XYZ")
        buf = io.BytesIO()
        wb.save(buf)
        self.client.force_login(self.sup.user)
        upload = SimpleUploadedFile("f50.xlsx", buf.getvalue())
        self.client.post(reverse("schedules:excel_import", args=[s.pk]), {"file": upload})
        row = s.rows.get(employee=self.alice)
        self.assertEqual(rules.row_totals(list(row.cells.select_related("shift"))), (36, 3))
        self.assertIsNone(row.cells.get(date=date(YEAR, MONTH, 4)).shift)
        # the downloaded template opens and carries the codes sheet
        resp = self.client.get(reverse("schedules:excel_template", args=[s.pk]))
        self.assertIn("Shift codes", load_workbook(io.BytesIO(resp.content)).sheetnames)


class PrintoutTests(Base):
    def test_layout_is_one_landscape_long_bond_page_with_real_names(self):
        from .form_f50 import fill_schedule_form

        s = self.new_schedule()
        self.route_to(s, S.APPROVED)
        ws = fill_schedule_form(s).active
        self.assertEqual(ws.page_setup.orientation, "landscape")
        self.assertEqual(ws.page_setup.paperSize, 14)
        self.assertEqual((ws.page_setup.fitToWidth, ws.page_setup.fitToHeight), (1, 1))
        text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
        self.assertIn("BDH-ADM-AO-01F50", text)
        self.assertIn("Revision: 2", text)
        self.assertIn(self.coh.full_name, text)
        self.assertIn("DIGITALLY APPROVED", text)
        self.assertIn("RECORDED BY HR", text)
        self.assertIn("PENDING", text)  # recorded step not reached yet

    @skipUnless(shutil.which("soffice") or shutil.which("libreoffice"), "LibreOffice not installed")
    def test_pdf_is_one_page(self):
        from .form_f50 import render_pdf

        s = self.new_schedule()
        for i in range(25):  # a big section still fits one page
            services.add_row(s, emp(f"schbig{i}", self.lab))
        pdf = render_pdf(s)
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertEqual(len(re.findall(rb"/Type\s*/Page(?![s\w])", pdf)), 1)
