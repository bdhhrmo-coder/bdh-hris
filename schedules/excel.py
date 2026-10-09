"""Excel template and import for a monthly duty schedule, laid out like
Form BDH-ADM-AO-01F50 (Batch 5). One file per schedule: Employee ID, Name,
Designation, one column per day (shift code from a dropdown), and Total
Hours / Total Days worked out by Excel from the 'Shift codes' sheet."""

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from django.core.exceptions import ValidationError

from . import services
from .models import ScheduleCell, ShiftCode

HEADER_ROW = 6
FIRST_DATA_ROW = 8
NAVY = PatternFill("solid", fgColor="0B2C4D")
WHITE_BOLD = Font(bold=True, color="FFFFFF", name="Arial", size=9)
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_template(schedule):
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    codes = list(ShiftCode.objects.filter(is_active=True))
    days = schedule.days
    first_day_col, last_day_col = 4, 3 + len(days)
    hours_col, days_col = last_day_col + 1, last_day_col + 2

    ws["A1"] = "PERSONNEL WORK SCHEDULE"
    ws["A1"].font = Font(bold=True, size=14, name="Arial")
    ws["A2"] = "BDH-ADM-AO-01F50  ·  Revision 2"
    ws["A3"] = f"Department/Section: {schedule.area_full_name}"
    ws["A4"] = f"Month/Year: {schedule.month_label}"
    for c in ("A2", "A3", "A4"):
        ws[c].font = Font(name="Arial", size=10, bold=c != "A2")

    heads = ["EMPLOYEE ID", "NAME (Last name, First name)", "DESIGNATION"]
    for i, h in enumerate(heads, start=1):
        ws.cell(row=HEADER_ROW, column=i, value=h)
        ws.merge_cells(start_row=HEADER_ROW, start_column=i, end_row=HEADER_ROW + 1, end_column=i)
    for i, d in enumerate(days):
        ws.cell(row=HEADER_ROW, column=first_day_col + i, value=d.day)
        ws.cell(row=HEADER_ROW + 1, column=first_day_col + i, value=d.strftime("%a"))
    for col, label in ((hours_col, "TOTAL HOURS"), (days_col, "TOTAL DAYS")):
        ws.cell(row=HEADER_ROW, column=col, value=label)
        ws.merge_cells(start_row=HEADER_ROW, start_column=col, end_row=HEADER_ROW + 1, end_column=col)
    for r in (HEADER_ROW, HEADER_ROW + 1):
        for c in range(1, days_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill, cell.font, cell.border = NAVY, WHITE_BOLD, BOX
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ref = wb.create_sheet("Shift codes")
    ref.append(["CODE", "PAID HOURS", "DUTY DAY (1/0)", "DESCRIPTION"])
    for c in codes:
        ref.append([c.code, float(c.paid_hours), 1 if c.is_duty else 0, c.description])
    n = len(codes) + 1
    codes_rng, hours_rng, duty_rng = f"'Shift codes'!$A$2:$A${n}", f"'Shift codes'!$B$2:$B${n}", f"'Shift codes'!$C$2:$C${n}"

    dv = DataValidation(type="list", formula1=codes_rng, allow_blank=True)
    dv.error = "Pick a shift code from the list (see the 'Shift codes' sheet)."
    ws.add_data_validation(dv)

    grid = {(c.row_id, c.date): c for c in
            ScheduleCell.objects.filter(row__schedule=schedule).select_related("shift")}
    rows = list(schedule.rows.select_related("employee"))
    for i, row in enumerate(rows):
        r = FIRST_DATA_ROW + i
        ws.cell(row=r, column=1, value=row.employee.employee_id).number_format = "@"
        ws.cell(row=r, column=2, value=f"{row.employee.surname.upper()}, {row.employee.first_name}")
        ws.cell(row=r, column=3, value=row.designation)
        for j, d in enumerate(days):
            cell = grid.get((row.pk, d))
            ws.cell(row=r, column=first_day_col + j, value=cell.shift.code if cell and cell.shift_id else None)
        rng = f"{get_column_letter(first_day_col)}{r}:{get_column_letter(last_day_col)}{r}"
        ws.cell(row=r, column=hours_col, value=f"=SUMPRODUCT(COUNTIF({rng},{codes_rng}),{hours_rng})")
        ws.cell(row=r, column=days_col, value=f"=SUMPRODUCT(COUNTIF({rng},{codes_rng}),{duty_rng})")
        dv.add(rng)
        for c in range(1, days_col + 1):
            ws.cell(row=r, column=c).border = BOX
            if c >= first_day_col:
                ws.cell(row=r, column=c).alignment = Alignment(horizontal="center")
    end = FIRST_DATA_ROW + len(rows) + 1
    ws.cell(row=end, column=1, value="NOTES/ REMARKS:").font = Font(bold=True, name="Arial", size=9)
    ws.cell(row=end, column=2, value=schedule.notes)
    ws.cell(row=end + 2, column=1, value=(
        "How to use: type or pick a shift code in each day cell (blank = not scheduled). Do not change the "
        "Employee ID column or the day numbers. Save, then upload this file on the schedule's edit page."
    )).font = Font(italic=True, size=8, name="Arial")

    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 16
    for c in range(first_day_col, last_day_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 6.5
    ws.column_dimensions[get_column_letter(hours_col)].width = 9
    ws.column_dimensions[get_column_letter(days_col)].width = 9
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=first_day_col)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 14
    return wb


def import_workbook(schedule, upload):
    """Fill the schedule's draft cells from a filled template. Returns
    (cells filled, [problems])."""
    try:
        wb = load_workbook(upload, data_only=True)
    except Exception:
        raise ValidationError("The file could not be read as an Excel workbook.")
    ws = wb["Schedule"] if "Schedule" in wb.sheetnames else wb.active
    header_row = None
    for r in range(1, 20):
        if str(ws.cell(row=r, column=1).value or "").strip().upper() == "EMPLOYEE ID":
            header_row = r
            break
    if header_row is None:
        raise ValidationError("This doesn't look like the schedule template (no 'EMPLOYEE ID' column).")
    day_cols = {}
    for c in range(4, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if isinstance(v, (int, float)) and 1 <= int(v) <= 31:
            day_cols[c] = int(v)
    month_days = {d.day: d for d in schedule.days}
    codes = {c.code.upper(): c.pk for c in ShiftCode.objects.filter(is_active=True)}
    rows = {r.employee.employee_id.upper(): r for r in schedule.rows.select_related("employee")}
    values, problems = {}, []
    for r in range(header_row + 2, ws.max_row + 1):
        emp_id = str(ws.cell(row=r, column=1).value or "").strip().upper()
        if not emp_id or emp_id.startswith("NOTES"):
            continue
        row = rows.get(emp_id)
        if row is None:
            problems.append(f"{emp_id} is not on this schedule (add the employee on the page first)")
            continue
        for c, day in day_cols.items():
            if day not in month_days:
                continue
            raw = str(ws.cell(row=r, column=c).value or "").strip().upper()
            if raw and raw not in codes:
                problems.append(f"{emp_id} day {day}: unknown code '{raw}' (left blank)")
            values[(row.pk, month_days[day])] = codes.get(raw)
    services.save_grid(schedule, values)
    return sum(1 for v in values.values() if v), problems
