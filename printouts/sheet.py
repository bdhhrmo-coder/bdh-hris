"""
Draws BDH printable forms into an openpyxl worksheet, which
leave/pdf_convert.workbook_to_pdf() then turns into a PDF.

FormBuilder keeps a running row counter so each form module reads top to
bottom like the paper form it produces:

    f = FormBuilder("APPLICATION FOR CTO", "BDH-ADM-HR-01F04-B")
    f.section("1.  APPLICANT INFORMATION")
    f.field("Name:", employee.full_name)
    ...
    f.stamps(stamps, stopped)
    f.footer(reference_no, printed_by)
    return f.wb

The page is a 7-column grid (A-G). The letterhead matches the COSP Leave
form (BDH-ADM-HR-01F04-A) so the whole 01F04 series looks the same.
"""

from pathlib import Path

import openpyxl
from django.utils import timezone
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .stamps import actor_name

ASSETS = Path(__file__).parent / "assets"
BDH_LOGO = ASSETS / "bdh_logo.png"
PALAWAN_SEAL = ASSETS / "palawan_seal.jpg"

FONT = "Arial"
THIN = Side(style="thin")
BOTTOM = Border(bottom=THIN)
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
SECTION_FILL = PatternFill("solid", fgColor="DDDDDD")

# Stamp colours: dark green for done, grey for pending, dark red for stopped.
DONE_COLOR = "1B5E20"
PENDING_COLOR = "808080"
STOPPED_COLOR = "B71C1C"

COLUMN_WIDTHS = [22, 11, 11, 18, 9, 9, 4]   # A..G, same as the COSP form
LEFT_BLOCK = ("A", "C")
RIGHT_BLOCK = ("D", "G")
FOLIO = 14  # openpyxl paperSize code for 8.5 x 13 in ("long bond"), as on F10


def merge_set(ws, cell_range, value, *, bold=False, italic=False, size=11, align="left",
              border=None, color=None, valign="center", wrap=True):
    ws.merge_cells(cell_range)
    cell = ws[cell_range.split(":")[0]]
    cell.value = value
    cell.font = Font(bold=bold, italic=italic, size=size, name=FONT, color=color)
    cell.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
    if border:
        for row in ws[cell_range]:
            for c in row:
                c.border = border


def box_outline(ws, top, bottom, left_col, right_col, color="000000", style="thin"):
    """Draw a rectangle around a block of cells (merged or not)."""
    side = Side(style=style, color=color)
    cols = [chr(c) for c in range(ord(left_col), ord(right_col) + 1)]
    for r in range(top, bottom + 1):
        for col in cols:
            cell = ws[f"{col}{r}"]
            cell.border = Border(
                top=side if r == top else cell.border.top,
                bottom=side if r == bottom else cell.border.bottom,
                left=side if col == left_col else cell.border.left,
                right=side if col == right_col else cell.border.right,
            )


class FormBuilder:
    def __init__(self, title, form_code, *, subtitle="", revision=None, paper=FOLIO, column_widths=COLUMN_WIDTHS):
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.ws.title = "Form"
        self.ws.sheet_view.showGridLines = False
        for col, width in zip("ABCDEFG", column_widths):
            self.ws.column_dimensions[col].width = width
        ps = self.ws.page_setup
        ps.orientation = "portrait"
        ps.paperSize = paper
        ps.fitToWidth = 1
        ps.fitToHeight = 0
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.ws.page_margins.left = self.ws.page_margins.right = 0.5
        self.ws.page_margins.top = self.ws.page_margins.bottom = 0.5
        self.row = 1
        self._letterhead(title, form_code, subtitle, revision)

    # -- layout pieces ----------------------------------------------------

    def _letterhead(self, title, form_code, subtitle, revision):
        ws = self.ws
        ws["G1"] = f"Form No. {form_code}"
        ws["G1"].font = Font(size=8, italic=True, name=FONT)
        ws["G1"].alignment = Alignment(horizontal="right")
        if revision is not None:
            ws["G2"] = f"Revision: {revision}"
            ws["G2"].font = Font(size=8, italic=True, name=FONT)
            ws["G2"].alignment = Alignment(horizontal="right")
        merge_set(ws, "A2:F2" if revision is not None else "A2:G2", "Republic of the Philippines", italic=True, align="center")
        merge_set(ws, "A3:G3", "PROVINCIAL GOVERNMENT OF PALAWAN", bold=True, size=13, align="center")
        merge_set(ws, "A4:G4", "BATARAZA DISTRICT HOSPITAL", bold=True, size=13, align="center")
        merge_set(ws, "A5:G5", "Barangay Marangas, Bataraza, Palawan", italic=True, size=9, align="center")
        merge_set(ws, "A7:G7", title, bold=True, size=15, align="center")
        if subtitle:
            merge_set(ws, "A8:G8", subtitle, italic=True, size=9, align="center")
        if BDH_LOGO.exists():
            logo = XLImage(str(BDH_LOGO))
            logo.width = logo.height = 70
            ws.add_image(logo, "A1")
        self.row = 10

    def gap(self, n=1):
        self.row += n

    def section(self, text):
        r = self.row
        merge_set(self.ws, f"A{r}:G{r}", text, bold=True, size=11)
        for c in self.ws[f"A{r}:G{r}"][0]:
            c.fill = SECTION_FILL
        self.row += 1

    def field(self, label, value, *, height=None):
        r = self.row
        self.ws[f"A{r}"] = label
        self.ws[f"A{r}"].font = Font(bold=True, size=10, name=FONT)
        self.ws[f"A{r}"].alignment = Alignment(vertical="center")
        merge_set(self.ws, f"B{r}:G{r}", "" if value is None else str(value), size=11, border=BOTTOM)
        if height:
            self.ws.row_dimensions[r].height = height
        self.row += 1

    def table(self, headers, rows, spans):
        """headers/rows are lists of cell texts; spans is a list of
        (first_col, last_col) per column, e.g. [("A","B"), ("C","D")]."""
        def put(values, bold):
            r = self.row
            for text, (c1, c2) in zip(values, spans):
                rng = f"{c1}{r}:{c2}{r}" if c1 != c2 else f"{c1}{r}"
                if c1 != c2:
                    merge_set(self.ws, rng, text, bold=bold, size=9 if bold else 10, align="center", border=BOX)
                else:
                    cell = self.ws[rng]
                    cell.value = text
                    cell.font = Font(bold=bold, size=9 if bold else 10, name=FONT)
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = BOX
            self.row += 1
        put(headers, True)
        for values in rows:
            put(values, False)

    def note(self, text, *, size=8, italic=True, height=None):
        r = self.row
        merge_set(self.ws, f"A{r}:G{r}", text, italic=italic, size=size)
        if height:
            self.ws.row_dimensions[r].height = height
        self.row += 1

    # -- digital stamps ----------------------------------------------------

    def _stamp_block(self, top, cols, stamp):
        c1, c2 = cols
        ws = self.ws
        if stamp.state == "done":
            color, headline = DONE_COLOR, f"✔ {stamp.text}"
        else:
            color, headline = PENDING_COLOR, stamp.text
        merge_set(ws, f"{c1}{top}:{c2}{top}", stamp.role_label.upper(), bold=True, size=8, align="center", color="333333")
        for c in ws[f"{c1}{top}:{c2}{top}"][0]:
            c.fill = SECTION_FILL
        merge_set(ws, f"{c1}{top + 1}:{c2}{top + 1}", headline, bold=stamp.state == "done",
                  italic=stamp.state != "done", size=10, align="center", color=color)
        merge_set(ws, f"{c1}{top + 2}:{c2}{top + 2}", stamp.name or "—", bold=True, size=10, align="center")
        merge_set(ws, f"{c1}{top + 3}:{c2}{top + 3}", stamp.position, italic=True, size=8, align="center")
        merge_set(ws, f"{c1}{top + 4}:{c2}{top + 4}", stamp.when, size=9, align="center")
        box_outline(ws, top, top + 4, c1, c2, color=color if stamp.state == "done" else "999999",
                    style="medium" if stamp.state == "done" else "thin")

    def stamps(self, stamps, stopped=None):
        """Two stamps per row; then a full-width red banner if the request
        was returned/rejected/declined/cancelled."""
        for i in range(0, len(stamps), 2):
            top = self.row
            self._stamp_block(top, LEFT_BLOCK, stamps[i])
            if i + 1 < len(stamps):
                self._stamp_block(top, RIGHT_BLOCK, stamps[i + 1])
            self.row = top + 6  # 5 stamp rows + 1 spacer
        if stopped is not None:
            r = self.row
            text = f"✖ {stopped.text}  by  {stopped.name}"
            if stopped.position:
                text += f", {stopped.position}"
            text += f"  —  {stopped.when}"
            merge_set(self.ws, f"A{r}:G{r}", text, bold=True, size=10, align="center", color=STOPPED_COLOR)
            box_outline(self.ws, r, r + (1 if stopped.notes else 0), "A", "G", color=STOPPED_COLOR, style="medium")
            self.row += 1
            if stopped.notes:
                merge_set(self.ws, f"A{self.row}:G{self.row}", f"Remarks: {stopped.notes}", italic=True, size=9,
                          align="center", color=STOPPED_COLOR)
                box_outline(self.ws, r, self.row, "A", "G", color=STOPPED_COLOR, style="medium")
                self.row += 1
            self.row += 1

    def footer(self, reference_no, printed_by=None):
        printed = f"Printed {timezone.localtime():%b %d, %Y %I:%M %p}"
        if printed_by is not None:
            printed += f" by {actor_name(printed_by)}"
        self.note(
            "This form was processed electronically through the BDH Human Resource Information System (HRIS). "
            "Each stamp above records the user account, date and time of that action, as kept in the system's "
            "action log.",
            height=26,
        )
        self.note(f"Reference No.: {reference_no}     ·     {printed}", size=8)

    def finish(self):
        self.ws.print_area = f"A1:G{self.row - 1}"
        return self.wb
