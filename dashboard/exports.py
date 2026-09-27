"""
Excel/PDF export for the dashboard's two reports (CLAUDE.md §12: "Reports
exportable to both PDF and Excel"). Uses openpyxl and reportlab directly —
both already project dependencies — rather than the CSC-Form-6-style
"fill an .xlsx template, convert with LibreOffice" pipeline used for the
official leave forms: these reports have no fixed institutional template
to preserve, so a plain generated table is simpler and drops one moving
part (no LibreOffice subprocess).
"""

from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def _xlsx_response(filename, sheets):
    """sheets: list of (sheet_name, header_row, data_rows)."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, header, data in sheets:
        ws = wb.create_sheet(title=name[:31])
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in data:
            ws.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return response


def _pdf_response(filename, title, tables):
    """tables: list of (heading, header_row, data_rows)."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter))
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Title"]), Spacer(1, 0.2 * inch)]
    for heading, header, data in tables:
        elements.append(Paragraph(heading, styles["Heading2"]))
        table_data = [[str(cell) for cell in header]] + [[str(cell) for cell in row] for row in data]
        table = Table(table_data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a3e72")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ]
            )
        )
        elements.append(table)
        elements.append(Spacer(1, 0.3 * inch))
    doc.build(elements)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'inline; filename="{filename}.pdf"'
    return response


def export_leave_balance_report(leave_types, rows, fmt):
    header = ["Employee ID", "Name", "Section(s)"] + [lt.name for lt in leave_types]
    data = []
    for row in rows:
        employee = row["employee"]
        sections = ", ".join(section.name for section in employee.sections.all())
        values = ["" if balance is None else balance for balance in row["balances"]]
        data.append([employee.employee_id, str(employee), sections, *values])

    if fmt == "xlsx":
        return _xlsx_response("leave_balance_report", [("Leave Balances", header, data)])
    return _pdf_response(
        "leave_balance_report", "BDH HRIS &mdash; Leave Balance Report", [("Leave Balances", header, data)]
    )


def export_application_summary_report(period_label, type_rows, totals, department_rows, fmt):
    type_header = ["Request Type", "Pending", "Approved", "Not Approved"]
    type_data = [[row["label"], row["pending"], row["approved"], row["not_approved"]] for row in type_rows]
    type_data.append(["TOTAL", totals["pending"], totals["approved"], totals["not_approved"]])

    dept_header = ["Section", "Headcount", "Pending", "Approved", "Not Approved"]
    dept_data = [
        [row["section"].name, row["headcount"], row["pending"], row["approved"], row["not_approved"]]
        for row in department_rows
    ]

    if fmt == "xlsx":
        return _xlsx_response(
            "application_summary_report",
            [("By Request Type", type_header, type_data), ("By Department", dept_header, dept_data)],
        )
    return _pdf_response(
        "application_summary_report",
        f"BDH HRIS &mdash; Application Summary ({period_label})",
        [("By Request Type", type_header, type_data), ("By Department", dept_header, dept_data)],
    )
