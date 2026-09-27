"""
Fills the actual CSC Form No. 6 (2020 revision, BDH copy — signatories
already printed on it: PGDH Rolando B. Buñi, Chief of Hospital Dr. Bon
Karlo C. Ladia, Governor Amy Roa Alvarez) rather than redrawing the form
from scratch, per the project's institutional-template-fidelity rule.

Only the blank answer cells are touched; everything else in the template
(headers, legal citations, signatories, instructions pages) is left as-is.

Column A next to each leave-type label (rows 16-28) is that leave type's
checkbox on the real form. COSP Leave prints on its own custom form (see
cosp_leave_form.py) — it never uses this one. Wellness and Emergency Leave
have no row of their own on the official CSC form, so they're marked
against the "Others:" line with the leave type name written in.

Commutation (6.D) isn't collected anywhere in this system yet (it's a
payroll-adjacent concept, out of Phase 4's scope) — this always marks
"Not Requested" as a placeholder default. Flagged here, not hidden.
"""

import subprocess
import tempfile
from copy import copy
from pathlib import Path

import openpyxl

TEMPLATE_PATH = Path(__file__).parent / "form_templates" / "CSC_Form_6_blank.xlsx"

# Position titles are free text and can run long ("Nursing Attendant / OIC,
# ER Unit"); the official form only gives that field ~26 characters of
# width before it would overlap "5. SALARY" next to it. Rather than
# truncating someone's job title, shrink the font just enough to fit.
POSITION_FIELD_MAX_CHARS = 26
POSITION_FIELD_MIN_FONT_SIZE = 8


def _write_shrink_to_fit(ws, coordinate, text, max_chars, min_size):
    cell = ws[coordinate]
    cell.value = text
    if text and len(text) > max_chars:
        base_size = cell.font.sz or 11
        fitted_size = max(min_size, base_size * max_chars / len(text))
        new_font = copy(cell.font)
        new_font.sz = fitted_size
        cell.font = new_font

# leave_type.code -> the row of its checkbox in column A on the real form.
LEAVE_TYPE_CHECKBOX_ROW = {
    "VL": 16,
    "MANDATORY": 17,
    "SL": 18,
    "MATERNITY": 19,
    "PATERNITY": 20,
    "SPL": 21,
    "SOLO_PARENT": 22,
    "STUDY": 23,
    "VAWC": 24,
    "REHAB": 25,
    "SPECIAL_WOMEN": 26,
    "CALAMITY": 27,
    "ADOPTION": 28,
}


def fill_csc_form6(application):
    """Returns an openpyxl Workbook with `application`'s data filled in."""
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb["blank"]
    employee = application.employee

    ws["K11"] = employee.surname
    ws["T11"] = employee.first_name
    ws["Z11"] = employee.middle_name

    ws["F12"] = application.submitted_at.strftime("%m/%d/%Y")  # narrow cell (F12:J12) — long month names overflow
    _write_shrink_to_fit(ws, "O12", employee.position, POSITION_FIELD_MAX_CHARS, POSITION_FIELD_MIN_FONT_SIZE)
    ws["AA12"] = employee.salary_grade

    code = application.leave_type.code
    row = LEAVE_TYPE_CHECKBOX_ROW.get(code)
    if row:
        ws.cell(row=row, column=1, value="X")
    else:
        ws["A30"] = "X"
        ws["B31"] = application.leave_type.name  # "Others:" (B30) is answered on the ruled line below it (merged B31:O31)

    ws["B34"] = str(application.number_of_days)
    ws["B36"] = f"{application.start_date:%B %d, %Y} to {application.end_date:%B %d, %Y}"
    ws["R35"] = "X"  # Commutation: Not Requested (placeholder — not modeled yet)

    return wb


def render_pdf(application) -> bytes:
    """Fills the form and converts it to PDF bytes via headless LibreOffice."""
    wb = fill_csc_form6(application)
    with tempfile.TemporaryDirectory() as tmpdir:
        xlsx_path = Path(tmpdir) / "form.xlsx"
        wb.save(xlsx_path)
        result = subprocess.run(
            [
                "libreoffice", "--headless", "--norestore",
                "--convert-to", "pdf", "--outdir", tmpdir, str(xlsx_path),
            ],
            capture_output=True, timeout=60,
        )
        pdf_path = Path(tmpdir) / "form.pdf"
        if not pdf_path.exists():
            raise RuntimeError(f"CSC Form 6 PDF conversion failed: {result.stderr.decode(errors='replace')}")
        return pdf_path.read_bytes()
