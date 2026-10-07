"""
Printable Missed Log Justification Form - BDH-ADM-AO-01F10, Revision 2
(project owner, 2026-10-06), for FORMAL attendance corrections.

Follows the Revision 1 paper slip: two logos and the form title/code at
the top; Name, Unit/Section, Date Filed; the date and log times; the
reason checkboxes with * (validated by ICTU) and ** (validated by HR);
then Requested / Validated* / Validated** / Approved.

What Revision 2 changes from the paper Revision 1:
  - One time-in / time-out pair per date. Revision 1 has AM IN/OUT and
    PM IN/OUT columns, but the HRIS keeps one pair per day (confirmed
    2026-09-27), so the printout shows Date / IN / OUT.
  - Paperless by default: each block is a "digitally ..." stamp with the
    person's name, position, date and time from the request's action log,
    instead of a signature line. The validator column that doesn't apply
    to the chosen reason prints NOT APPLICABLE.
  - The approver is whoever actually approved it in the system (an
    Administrative Officer), not a name printed in advance.
"""

from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font

from leave.pdf_convert import workbook_to_pdf
from printouts.sheet import BDH_LOGO, BOTTOM, BOX, FONT, PALAWAN_SEAL, FormBuilder, merge_set
from printouts.stamps import StepDef, Stamp, actor_name, compute_stamps, format_when

from .models import AttendanceCorrectionRequest

FORM_CODE = "BDH-ADM-AO-01F10"
REVISION = 2
# A:C and D:G are the two stamp columns, so keep them about equal.
COLUMN_WIDTHS = [11, 11, 22, 14, 10, 10, 10]

ROUTING = [
    StepDef("Requested (Employee)", ("submit", "resubmit"), "DIGITALLY REQUESTED"),
    StepDef("Validated* (ICTU Staff)", "validate", "DIGITALLY VALIDATED"),
    # "process" = HR's step on the old Supervisor -> HR -> AO chain, for
    # requests filed before Revision 2.
    StepDef("Validated** (HR Staff)", ("validate", "process"), "DIGITALLY VALIDATED"),
    StepDef("Approved (Administrative Officer)", "approve", "DIGITALLY APPROVED"),
]


def reference_no(correction):
    return f"MLJ-{correction.pk:06d}"


class F10Builder(FormBuilder):
    """Same page and stamps as the 01F04 forms, with the F10 header."""

    def _letterhead(self, title, form_code, subtitle, revision):
        ws = self.ws
        for path, anchor in ((PALAWAN_SEAL, "A1"), (BDH_LOGO, "B1")):
            if path.exists():
                img = XLImage(str(path))
                img.width = img.height = 58
                ws.add_image(img, anchor)
        merge_set(ws, "D1:G1", title, bold=True, size=12, align="right")
        merge_set(ws, "D2:G2", form_code, size=9, align="right")
        merge_set(ws, "D3:G3", f"Revision: {revision}", size=9, align="right")
        merge_set(ws, "D4:G4", "Bataraza District Hospital", italic=True, size=8, align="right")
        self.row = 6


def _line(f, label, value, value_cols="B:G"):
    r = f.row
    f.ws[f"A{r}"] = label
    f.ws[f"A{r}"].font = Font(bold=True, size=10, name=FONT)
    c1, c2 = value_cols.split(":")
    merge_set(f.ws, f"{c1}{r}:{c2}{r}", value, size=11, border=BOTTOM)
    f.row += 1


def fill_correction_form(correction, printed_by=None):
    employee = correction.employee
    f = F10Builder("MISSED LOG JUSTIFICATION FORM", FORM_CODE, revision=REVISION, column_widths=COLUMN_WIDTHS)
    ws = f.ws

    _line(f, "Name:", employee.full_name)
    r = f.row
    ws[f"A{r}"] = "Unit/Section:"
    ws[f"A{r}"].font = Font(bold=True, size=10, name=FONT)
    merge_set(ws, f"B{r}:C{r}", ", ".join(s.name for s in employee.sections.all()), size=10, border=BOTTOM)
    ws[f"D{r}"] = "Date Filed:"
    ws[f"D{r}"].font = Font(bold=True, size=10, name=FONT)
    ws[f"D{r}"].alignment = Alignment(horizontal="right")
    merge_set(ws, f"E{r}:G{r}", f"{correction.submitted_at:%m/%d/%Y}", size=11, border=BOTTOM)
    f.row += 2

    lines = correction.sorted_lines
    rows = []
    for line in lines:
        if line.is_absent:
            t_in = t_out = "Absent"
        else:
            t_in = f"{line.time_in:%I:%M %p}" if line.time_in else "—"
            t_out = f"{line.time_out:%I:%M %p}" if line.time_out else "—"
            if line.overnight:
                t_out += " (+1 day)"
        reason = line.get_reason_category_display() if line.reason_category else ""
        if line.reason:
            reason = f"{reason}: {line.reason}" if reason else line.reason
        rows.append([f"{line.date:%b %d, %Y}", t_in, t_out, reason])
    f.table(["Date", "IN", "OUT", "Reason"], rows, [("A", "B"), ("C", "C"), ("D", "D"), ("E", "G")])
    for i, line in enumerate(lines):
        r = f.row - len(lines) + i
        ws.row_dimensions[r].height = 30 if (len(rows[i][3]) > 28 or line.overnight) else 20
        ws[f"D{r}"].alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        for col in "EFG":
            ws[f"{col}{r}"].alignment = Alignment(wrap_text=True, vertical="center", horizontal="left", indent=1)
    f.gap()

    ws[f"A{f.row}"] = "Reason/s:"
    ws[f"A{f.row}"].font = Font(bold=True, size=10, name=FONT)
    chosen = {line.reason_category for line in lines}
    stars = {code: ("*" if code in AttendanceCorrectionRequest.ICTU_VALIDATED_REASONS else "**")
             for code, _ in AttendanceCorrectionRequest.REASON_CATEGORY_CHOICES}
    for code, label in AttendanceCorrectionRequest.REASON_CATEGORY_CHOICES:
        mark = "☒" if code in chosen else "☐"
        merge_set(ws, f"B{f.row}:G{f.row}", f"{mark}  {label}{stars[code]}", size=10, bold=code in chosen)
        f.row += 1
    f.note("*  Validated by ICTU Staff        **  Validated by HR Staff        (details per date in the table above)",
           size=8)
    f.gap()

    returns = list(correction.actions.filter(action="return").select_related("acted_by__employee").order_by("acted_at"))
    if returns:
        f.section("RETURN HISTORY")
        f.table(["Returned on", "By", "Remark"],
                [[format_when(a.acted_at), actor_name(a.acted_by), a.notes] for a in returns],
                [("A", "B"), ("C", "C"), ("D", "G")])
        f.gap()

    # Stamps show the current round only (after the latest resubmission).
    stamps, stopped = compute_stamps(correction.current_cycle_actions(), ROUTING)
    # Only one validator applies, depending on the reason chosen.
    not_applicable = 2 if correction.validated_by_ictu else 1
    stamps[not_applicable] = Stamp(role_label=stamps[not_applicable].role_label, state="pending",
                                   text="NOT APPLICABLE")
    f.stamps(stamps, stopped)
    f.footer(reference_no(correction), printed_by)
    return f.finish()


def render_pdf(correction, printed_by=None) -> bytes:
    return workbook_to_pdf(fill_correction_form(correction, printed_by), "Missed Log Justification Form")
