"""The downloadable Excel template (one file, both imports) and the
temporary-password slips."""

from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from employees.models import Employee
from orgstructure.models import Section

from . import balance_import, employee_import

HEADER_FILL = PatternFill("solid", fgColor="0B2C4D")
REQUIRED_FILL = PatternFill("solid", fgColor="7A1F1F")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)
THIN = Side(style="thin", color="999999")

INSTRUCTIONS = [
    ("BDH HRIS — Data Import Template", True),
    ("", False),
    ("1. EMPLOYEES sheet — one row per employee.", True),
    ("• Columns marked * are required for a NEW employee. Red headers = required.", False),
    ("• If the Employee ID already exists, that employee is UPDATED. Blank cells never erase existing data.", False),
    ("• Dates: type as YYYY-MM-DD (e.g. 2024-01-15). Excel dates are fine too.", False),
    ("• Section(s) and Unit(s): copy the exact names from the 'Sections and Units' sheet. Separate several with ;", False),
    ("• Additional Role(s): only for an HR Administrator running the import. Separate several with ;", False),
    ("   Allowed: Supervisor; HR Processor; HR Administrator; Administrative Officer; Chief of Hospital; ICTU Staff.", False),
    ("   Everyone automatically gets the Employee role. A Supervisor supervises the section(s) on their row.", False),
    ("   System Administrator cannot be given by import. Imports only ADD roles; they never remove one.", False),
    ("• Username: leave blank to use the Employee ID. Staff can log in with either.", False),
    ("• New login accounts get a temporary password. After the import, print the password slips and hand", False),
    ("   them out personally. Each person must set their own password at first login.", False),
    ("", False),
    ("2. OPENING BALANCES sheet — import AFTER the employees.", True),
    ("• Enter each balance exactly as it stands on the employee's leave card on the As-of Date.", False),
    ("• Use the date the leave card is updated to (e.g. the last day of the month already credited).", False),
    ("• Leave a cell blank to leave that balance as it is. Days may have up to 2 decimals (e.g. 12.38).", False),
    ("• Wellness Leave = days still remaining this year (maximum 5).", False),
    ("• The system records the difference as a 'Manual adjustment' dated the As-of Date, so the", False),
    ("   balance then matches the leave card. Importing the same file twice adds nothing.", False),
    ("", False),
    ("3. Upload the file on HR Tools → Data Import. You will see a preview; nothing is saved until you", False),
    ("   click Confirm, and nothing is saved if any row has an error.", False),
]


def _header(ws, headers, required_keys=()):
    for col, (key, title) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = HEADER_FONT
        cell.fill = REQUIRED_FILL if key in required_keys else HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[cell.column_letter].width = max(14, min(32, len(title) + 4))
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"


def _list_validation(ws, col_letter, options, rows=1000):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    dv.error = "Please pick from the list."
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{rows}")


def build_template():
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = 110
    for i, (line, bold) in enumerate(INSTRUCTIONS, start=1):
        ws.cell(row=i, column=1, value=line).font = Font(bold=bold, name="Arial", size=12 if i == 1 else 10)

    # Employees
    ws = wb.create_sheet(employee_import.SHEET)
    cols = [(k, h) for k, h, _ in employee_import.COLUMNS]
    _header(ws, cols, {k for k, _, req in employee_import.COLUMNS if req})
    letters = {k: openpyxl.utils.get_column_letter(i) for i, (k, _) in enumerate(cols, start=1)}
    _list_validation(ws, letters["employment_status"], ["Regular", "COSP"])
    _list_validation(ws, letters["sex_at_birth"], ["Male", "Female"])
    _list_validation(ws, letters["civil_status"], [label for _, label in Employee.CIVIL_STATUS_CHOICES])
    _list_validation(ws, letters["shift_hours"], ["8", "12"])
    for key in employee_import.DATE_FIELDS:
        for r in range(2, 1001):
            ws[f"{letters[key]}{r}"].number_format = "yyyy-mm-dd"
    for key in ("employee_id", "sss_number", "pagibig_number", "philhealth_number", "tin_number", "gsis_number",
                "salary_grade", "item_plantilla_no", "telephone_mobile"):
        for r in range(2, 1001):
            ws[f"{letters[key]}{r}"].number_format = "@"  # keep leading zeros

    # Opening balances, pre-listed with every active employee already in the system
    ws = wb.create_sheet(balance_import.SHEET)
    _header(ws, balance_import.COLUMNS, {"employee_id", "as_of"})
    for r, emp in enumerate(Employee.objects.filter(is_active=True).order_by("surname", "first_name"), start=2):
        ws.cell(row=r, column=1, value=emp.employee_id).number_format = "@"
        ws.cell(row=r, column=2, value=emp.full_name)
    for r in range(2, 1001):
        ws[f"C{r}"].number_format = "yyyy-mm-dd"
    ws.column_dimensions["B"].width = 32

    # Reference list
    ws = wb.create_sheet("Sections and Units")
    ws["A1"], ws["B1"] = "Section (copy exactly)", "Units in this section"
    for c in ("A1", "B1"):
        ws[c].font, ws[c].fill = HEADER_FONT, HEADER_FILL
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 45, 60
    for r, section in enumerate(Section.objects.filter(is_active=True).order_by("name"), start=2):
        ws.cell(row=r, column=1, value=section.name)
        ws.cell(row=r, column=2, value="; ".join(u.name for u in section.units.filter(is_active=True)))
    return wb


def build_password_slips(credentials, login_url):
    """One cut-out slip per new account, three per page."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Password slips"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 60
    ws.page_setup.paperSize = 14  # 8.5 x 13 in
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    box = Border(top=Side(style="dashed"), bottom=Side(style="dashed"), left=Side(style="dashed"),
                 right=Side(style="dashed"))
    row = 1
    for i, c in enumerate(credentials):
        lines = [
            ("BDH HRIS — YOUR LOGIN (CONFIDENTIAL)", ""),
            ("Name:", c["name"]),
            ("Employee ID:", c["employee_id"]),
            ("Username:", c["username"]),
            ("Temporary password:", c["password"]),
            ("Log in at:", login_url),
            ("", "You will be asked to set your own password at first login. Do not share it."),
        ]
        top = row
        for label, value in lines:
            a, b = ws.cell(row=row, column=1, value=label), ws.cell(row=row, column=2, value=value)
            a.font = Font(bold=True, name="Arial", size=10 if row != top else 11)
            b.font = Font(name="Consolas" if label == "Temporary password:" else "Arial",
                          bold=label in ("Temporary password:", "Username:"), size=11 if label == "Temporary password:" else 10)
            row += 1
        for r in range(top, row):
            for col in (1, 2):
                cell = ws.cell(row=r, column=col)
                cell.border = Border(top=box.top if r == top else None, bottom=box.bottom if r == row - 1 else None,
                                     left=box.left if col == 1 else None, right=box.right if col == 2 else None)
        row += 2
        if (i + 1) % 3 == 0:
            ws.row_breaks.append(openpyxl.worksheet.pagebreak.Break(id=row - 1))
    ws.print_area = f"A1:B{max(row - 1, 1)}"
    ws["D1"] = f"Generated {date.today():%Y-%m-%d}"
    ws["D1"].font = Font(italic=True, size=8, color="999999")
    return wb
